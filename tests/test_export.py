"""Tests for Excel/PDF export and order_service validation (Phase 9)."""

from __future__ import annotations

import pytest

from app.export.excel_exporter import export_excel
from app.export.pdf_exporter import export_pdf
from app.models.product import Product, format_price, parse_price
from app.services.order_service import build_order, parse_quantity

from openpyxl import load_workbook


def _product(code: str, name: str = "Neki proizvod") -> Product:
    return Product(
        id=0,
        code=code,
        name=name,
        package="4",
        image_path=None,
        category="Kategorija",
        page_number=1,
        raw_text="",
    )


# ---------------------------------------------------------------- quantity


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("1", (1, None)),
        (" 5 ", (5, None)),
        ("10", (10, None)),
        ("", (None, "empty")),
        ("0", (None, "zero")),
        ("00", (None, "zero")),
        ("-5", (None, "negative")),
        ("3.5", (None, "decimal")),
        ("3,5", (None, "decimal")),
        ("abc", (None, "text")),
        ("-", (None, "text")),
    ],
)
def test_parse_quantity(raw: str, expected) -> None:
    assert parse_quantity(raw) == expected


def test_build_order_valid() -> None:
    order, errors = build_order([_product("12391")], {"12391": "12"})
    assert errors == []
    assert len(order.items) == 1
    assert order.items[0].quantity == 12


def test_build_order_rejects_invalid_and_drops_item() -> None:
    order, errors = build_order(
        [_product("12391"), _product("88888")],
        {"12391": "", "88888": "2"},
    )
    assert len(errors) == 1
    assert len(order.items) == 1
    assert order.items[0].product.code == "88888"


# ------------------------------------------------------------------ excel


def test_excel_export_writes_code_and_quantity_only(tmp_path) -> None:
    path = export_excel(
        [("12391", 12), ("M456", 3)], path=tmp_path / "out.xlsx"
    )
    workbook = load_workbook(path)
    sheet = workbook.active
    assert sheet.max_row == 3
    assert sheet["A1"].value == "Šifra"
    assert sheet["B1"].value == "Količina"
    assert sheet["A2"].value == "12391"
    assert sheet["B2"].value == 12
    assert sheet["A3"].value == "M456"
    assert sheet["B3"].value == 3
    assert sheet.max_column == 2  # no prices/names/images leaked


def test_excel_export_appends_approximate_total(tmp_path) -> None:
    path = export_excel(
        [("12391", 12, 1033.70), ("11951", 2, 249.31)],
        path=tmp_path / "total.xlsx",
    )
    sheet = load_workbook(path).active
    assert sheet.max_row == 4
    assert sheet["A4"].value == "UKUPNO (približno), RSD"
    assert sheet["B4"].value == pytest.approx(12 * 1033.70 + 2 * 249.31)


def test_excel_export_skips_total_when_no_prices(tmp_path) -> None:
    path = export_excel([("12391", 12, None)], path=tmp_path / "no.xlsx")
    sheet = load_workbook(path).active
    assert sheet.max_row == 2  # data only, no total row


# ---------------------------------------------------------------- pricing


def test_format_price_serbian() -> None:
    assert format_price(1033.7) == "1.033,70"
    assert format_price(12903.02) == "12.903,02"
    assert parse_price("1,033.70") == pytest.approx(1033.70)
    assert parse_price("249.31") == pytest.approx(249.31)
    assert parse_price("-") is None


def test_price_for_uses_correct_tier() -> None:
    product = Product(
        id=0,
        code="11951",
        name="X",
        tiers=[("1-5 kom", "249.31"), ("6-11 kom", "236.59"), ("12+ kom", "228.96")],
    )
    assert product.price_for(3) == pytest.approx(249.31)
    assert product.price_for(5) == pytest.approx(249.31)
    assert product.price_for(6) == pytest.approx(236.59)
    assert product.price_for(11) == pytest.approx(236.59)
    assert product.price_for(12) == pytest.approx(228.96)
    assert product.price_for(200) == pytest.approx(228.96)  # bulk -> last tier


def test_price_for_none_when_unpriced() -> None:
    priceless = Product(id=0, code="11296", name="X", tiers=[("1-3 kom", "-")])
    assert priceless.price_for(2) is None
    empty = Product(id=0, code="Y", name="X")
    assert empty.price_for(2) is None


# -------------------------------------------------------------------- pdf


def test_pdf_export_creates_file(tmp_path) -> None:
    path = export_pdf(
        [("12391", 12), ("M456", 3)], path=tmp_path / "out.pdf"
    )
    assert path.exists()
    assert path.stat().st_size > 1000

    import fitz  # PyMuPDF

    with fitz.open(path) as doc:
        text = "".join(page.get_text() for page in doc)

    assert "PORUDŽBINA" in text
    assert "Šifra" in text and "Količina" in text
    assert "M456" in text
    assert "3" in text


def test_pdf_export_puts_headers_in_two_columns(tmp_path) -> None:
    """Šifra and Količina must sit side by side in ONE header row."""
    path = export_pdf([("111", 2)], path=tmp_path / "headers.pdf")

    import fitz  # PyMuPDF

    with fitz.open(path) as doc:
        words = doc[0].get_text("words")

    header = [w for w in words if w[4] in ("Šifra", "Količina")]
    assert len(header) == 2
    # both header cells share the same baseline row (y0)
    assert abs(header[0][1] - header[1][1]) < 2
    # Šifra sits left of Količina
    assert header[0][0] < header[1][0]


def test_pdf_export_prints_approximate_total(tmp_path) -> None:
    path = export_pdf(
        [("12391", 12, 1033.70), ("11951", 2, 249.31)],
        path=tmp_path / "total.pdf",
    )
    import fitz  # PyMuPDF

    with fitz.open(path) as doc:
        text = "".join(page.get_text() for page in doc)

    assert "PRIBLIŽNA UKUPNA CENA (RSD)" in text
    assert "12.903,02" in text


def test_pdf_export_omits_total_when_no_prices(tmp_path) -> None:
    path = export_pdf([("12391", 12)], path=tmp_path / "plain.pdf")
    import fitz  # PyMuPDF

    with fitz.open(path) as doc:
        text = "".join(page.get_text() for page in doc)

    assert "PRIBLIŽNA UKUPNA CENA" not in text