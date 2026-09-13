"""Excel export (Phase 6): CODE + QUANTITY only, no prices, names or images."""

from __future__ import annotations

import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font

HEADERS = ("Šifra", "Količina")


def default_filename() -> str:
    return f"porudzbina_{datetime.date.today().isoformat()}.xlsx"


def export_excel(items: Iterable[tuple[str, int]], path: str | Path | None = None) -> Path:
    """Export (code, quantity) pairs to an .xlsx file.

    Returns the path that was written.
    """
    target = Path(path) if path else Path(default_filename())
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Porudžbina"

    sheet.cell(row=1, column=1, value=HEADERS[0]).font = Font(bold=True)
    sheet.cell(row=1, column=2, value=HEADERS[1]).font = Font(bold=True)

    for row_index, (code, quantity) in enumerate(items, start=2):
        sheet.cell(row=row_index, column=1, value=str(code))
        sheet.cell(row=row_index, column=2, value=int(quantity))

    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 12
    workbook.save(str(target))
    return target